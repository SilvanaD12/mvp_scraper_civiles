"""Exportacion profesional de resultados a Microsoft Excel."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


class ExportError(RuntimeError):
    """Error al generar el reporte Excel."""


class ExcelExporter:
    """Escribe hojas normalizadas y consolidadas con formato consistente."""

    HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
    HEADER_FONT = Font(color="FFFFFF", bold=True)

    def export_report(
        self,
        normalized: pd.DataFrame,
        consolidated: pd.DataFrame,
        destination: Path,
        title: str = "Reporte de costos civiles",
    ) -> Path:
        if normalized.empty and consolidated.empty:
            raise ExportError("No existen datos para exportar")
        destination.parent.mkdir(parents=True, exist_ok=True)
        try:
            with pd.ExcelWriter(destination, engine="openpyxl") as writer:
                normalized.to_excel(writer, sheet_name="Precios", index=False)
                consolidated.to_excel(writer, sheet_name="Consolidado", index=False)
                metadata = pd.DataFrame(
                    {"Campo": ["Reporte", "Registros normalizados", "Grupos consolidados"],
                     "Valor": [title, len(normalized), len(consolidated)]}
                )
                metadata.to_excel(writer, sheet_name="Resumen", index=False)
                for worksheet in writer.book.worksheets:
                    self._format_worksheet(worksheet)
        except (OSError, ValueError, ImportError) as exc:
            raise ExportError(f"No se pudo generar el archivo: {destination}") from exc
        return destination

    @classmethod
    def _format_worksheet(cls, worksheet: Worksheet) -> None:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_cells in worksheet.columns:
            length = max(len(str(cell.value or "")) for cell in column_cells)
            column_index = column_cells[0].column
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(length + 2, 12), 45)
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00'
